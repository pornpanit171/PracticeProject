#!/usr/bin/env python3
"""
make-template.py — สร้างไฟล์ tools/template.xlsx สำหรับกรอกโจทย์

รันใหม่เมื่อแก้ชื่อโดเมนใน data/domains.json:
    python tools/make-template.py

ต้องมี openpyxl:  pip install openpyxl
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
DOMAINS = json.loads((ROOT / "data" / "domains.json").read_text(encoding="utf-8"))

HEADERS = [
    ("id", 14, "เว้นว่างได้ ระบบจะรันเลขต่อให้เอง เช่น d1-0043"),
    ("domain", 9, "หมายเลขโดเมน 1-8 (จำเป็น)"),
    ("type", 10, "single หรือ multi เว้นว่าง = single"),
    ("q_en", 46, "คำถามภาษาอังกฤษ (จำเป็น)"),
    ("q_th", 46, "คำถามภาษาไทย"),
    ("a_en", 30, "ตัวเลือก A อังกฤษ"),
    ("a_th", 30, "ตัวเลือก A ไทย"),
    ("b_en", 30, "ตัวเลือก B อังกฤษ"),
    ("b_th", 30, "ตัวเลือก B ไทย"),
    ("c_en", 30, "ตัวเลือก C อังกฤษ"),
    ("c_th", 30, "ตัวเลือก C ไทย"),
    ("d_en", 30, "ตัวเลือก D อังกฤษ"),
    ("d_th", 30, "ตัวเลือก D ไทย"),
    ("e_en", 24, "ตัวเลือก E (ถ้ามี)"),
    ("e_th", 24, "ตัวเลือก E ไทย"),
    ("f_en", 24, "ตัวเลือก F (ถ้ามี)"),
    ("f_th", 24, "ตัวเลือก F ไทย"),
    ("answer", 12, 'เฉลย เช่น B หรือ "B,D" ถ้าตอบได้หลายข้อ (จำเป็น)'),
    ("explain_en", 52, "คำอธิบายภาษาอังกฤษ (ทำไมข้อนี้ถูก)"),
    ("explain_th", 52, "คำอธิบายภาษาไทย (ทำไมข้อนี้ถูก)"),
    ("insight_en", 52, "เพิ่มเติม/บทวิเคราะห์เชิงลึก อังกฤษ (ไม่บังคับ กด Alt+Enter ขึ้นบรรทัดใหม่ในเซลล์ได้)"),
    ("insight_th", 52, "เพิ่มเติม/บทวิเคราะห์เชิงลึก ไทย (ไม่บังคับ — ปกติกรอกแค่คอลัมน์นี้คอลัมน์เดียวก็พอ ระบบจะโชว์ให้ทั้งโหมด EN และ TH)"),
    ("tags", 20, "แท็ก คั่นด้วย ; เช่น risk;governance"),
    ("ref", 20, "อ้างอิง เช่น OSG ch.1"),
]

EXAMPLES = [
    {
        "id": "",
        "domain": 1,
        "type": "single",
        "q_en": "Which of the following BEST describes residual risk?",
        "q_th": 'ข้อใดอธิบาย "ความเสี่ยงคงเหลือ" ได้ดีที่สุด',
        "a_en": "The total risk before any control is applied",
        "a_th": "ความเสี่ยงทั้งหมดก่อนนำมาตรการควบคุมมาใช้",
        "b_en": "The risk that remains after controls are implemented",
        "b_th": "ความเสี่ยงที่ยังเหลืออยู่หลังใช้มาตรการควบคุมแล้ว",
        "c_en": "The risk transferred to an insurance provider",
        "c_th": "ความเสี่ยงที่โอนไปให้บริษัทประกันภัย",
        "d_en": "The risk from threats not yet identified",
        "d_th": "ความเสี่ยงจากภัยคุกคามที่ยังไม่ถูกระบุ",
        "answer": "B",
        "explain_en": "Whatever exposure is left after safeguards are applied is residual risk.",
        "explain_th": "ส่วนที่ยังเหลืออยู่หลังใส่มาตรการควบคุมแล้วคือความเสี่ยงคงเหลือ",
        "tags": "risk-management",
        "ref": "",
    },
    {
        "id": "",
        "domain": 1,
        "type": "multi",
        "q_en": "Besides confidentiality, which TWO principles make up the CIA triad?",
        "q_th": "นอกจาก Confidentiality แล้ว อีกสองหลักการใดที่ประกอบกันเป็น CIA triad",
        "a_en": "Integrity",
        "a_th": "Integrity (ความถูกต้องครบถ้วน)",
        "b_en": "Accountability",
        "b_th": "Accountability (ความรับผิดชอบที่ตรวจสอบได้)",
        "c_en": "Availability",
        "c_th": "Availability (ความพร้อมใช้งาน)",
        "d_en": "Authentication",
        "d_th": "Authentication (การพิสูจน์ตัวตน)",
        "answer": "A,C",
        "explain_en": "The triad is Confidentiality, Integrity and Availability.",
        "explain_th": "CIA triad ประกอบด้วย Confidentiality, Integrity และ Availability",
        "tags": "fundamentals",
        "ref": "",
    },
    {
        "id": "",
        "domain": 1,
        "type": "single",
        "q_en": (
            "Alyssa is responsible for her organization's security awareness program. "
            "She is concerned that changes in technology may make the content outdated.\n\n"
            "What control can she put in place to protect against this risk?"
        ),
        "q_th": (
            "Alyssa รับผิดชอบโปรแกรมสร้างตระหนักรู้ด้านความมั่นคงปลอดภัย (security awareness program) "
            "ขององค์กร เธอมีความกังวลว่าการเปลี่ยนแปลงทางเทคโนโลยีอาจทำให้เนื้อหาฝึกอบรมล้าสมัย\n\n"
            "มาตรการควบคุม (control) ใดที่เธอสามารถนำมาใช้เพื่อป้องกันความเสี่ยงนี้ได้?"
        ),
        "a_en": "Gamification",
        "a_th": "Gamification (การใช้เทคนิคเกมในการอบรม)",
        "b_en": "Computer-based training",
        "b_th": "Computer-based training (การฝึกอบรมผ่านคอมพิวเตอร์)",
        "c_en": "Content Reviews",
        "c_th": "Content Reviews (การทบทวนและตรวจสอบเนื้อหา)",
        "d_en": "Live training",
        "d_th": "Live training (การฝึกอบรมแบบสด/มีวิทยากร)",
        "answer": "C",
        "explain_en": (
            "Alyssa should use periodic content reviews to continually verify that the content in her "
            "program meets the organization's needs and is up-to-date based upon the evolving risk "
            "landscape. She may do this using a combination of computer-based training, live training, "
            "and gamification, but those techniques do not necessarily verify that the content is updated."
        ),
        "explain_th": (
            "Alyssa ควรใช้การทบทวนเนื้อหาเป็นระยะ (periodic content reviews) เพื่อตรวจสอบอย่างต่อเนื่องว่า"
            "เนื้อหาในโปรแกรมของเธอนั้นตอบสนองต่อความต้องการขององค์กร และทันสมัยอยู่เสมอโดยอ้างอิงตามสภาวะ"
            "ความเสี่ยงที่เปลี่ยนแปลงไป (evolving risk landscape)\n\n"
            "เธออาจใช้เทคนิคต่างๆ ร่วมกันได้ เช่น การฝึกอบรมผ่านคอมพิวเตอร์, การอบรมแบบสด และการใช้เทคนิคเกม "
            "แต่เทคนิคเหล่านั้นไม่ได้ช่วยตรวจสอบหรือยืนยันว่าเนื้อหาได้รับการอัปเดตแล้วแต่อย่างใด"
        ),
        # กรอกแค่ insight_th ก็พอ — ปล่อย insight_en ว่างไว้ ระบบจะโชว์อันนี้ให้ในโหมด EN ด้วยเช่นกัน (fallback อัตโนมัติ)
        "insight_en": "",
        "insight_th": (
            "1. Think Like a Manager / Fix the Process (คิดแบบผู้จัดการ เน้นที่กระบวนการ):\n"
            "     คนทำงานสายเทคนิคหรือคนทั่วไปมักจะมองหาเครื่องมือที่ฟังดูทันสมัย สะดุดตา (เช่น Gamification "
            "หรือ CBT) แต่ผู้จัดการความปลอดภัยที่ดีจะมองหา \"กระบวนการควบคุมและตรวจสอบคุณภาพ\" "
            "(Quality Assurance Process) ซึ่งในที่นี้คือการกำหนดตารางเวลาเพื่อ \"ทบทวนเนื้อหา\" (Content Reviews)\n\n"
            "2. Understand the Goal of Each Control (เข้าใจเป้าหมายที่แท้จริงของมาตรการ):\n"
            "     เป้าหมายของเทคนิคการสอน (A, B, D) มีไว้เพื่อเพิ่มการมีส่วนร่วม (Engagement) และการปรับเปลี่ยน"
            "พฤติกรรม (Behavior modification) ของพนักงาน\n"
            "     แต่เป้าหมายของ Content Reviews มีไว้เพื่อรับประกันความถูกต้อง ความสอดคล้อง และความสดใหม่ของ"
            "ข้อมูล (Information Integrity and Alignment)"
        ),
        "tags": "security-awareness;governance",
        "ref": "",
    },
]

HEAD_FILL = PatternFill("solid", fgColor="1F2937")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=10)
REQ_FILL = PatternFill("solid", fgColor="3B1220")

wb = Workbook()

# ── ชีต Questions ────────────────────────────────────────────
ws = wb.active
ws.title = "Questions"

for i, (name, width, note) in enumerate(HEADERS, start=1):
    cell = ws.cell(row=1, column=i, value=name)
    cell.fill = REQ_FILL if name in ("domain", "q_en", "answer") else HEAD_FILL
    cell.font = HEAD_FONT
    cell.alignment = Alignment(vertical="center")
    cell.comment = None
    ws.column_dimensions[get_column_letter(i)].width = width

for r, row in enumerate(EXAMPLES, start=2):
    for i, (name, _, _) in enumerate(HEADERS, start=1):
        ws.cell(row=r, column=i, value=row.get(name, ""))

ws.freeze_panes = "A2"
ws.row_dimensions[1].height = 22
for r in range(2, 400):
    for i in range(1, len(HEADERS) + 1):
        ws.cell(row=r, column=i).alignment = Alignment(vertical="top", wrap_text=True)

# dropdown ให้คอลัมน์ domain และ type
dv_domain = DataValidation(
    type="list",
    formula1='"' + ",".join(str(d["id"]) for d in DOMAINS) + '"',
    allow_blank=False,
    showErrorMessage=True,
    errorTitle="โดเมนไม่ถูกต้อง",
    error="เลือกหมายเลขโดเมนจากรายการเท่านั้น",
)
dv_type = DataValidation(
    type="list", formula1='"single,multi"', allow_blank=True, showErrorMessage=True
)
ws.add_data_validation(dv_domain)
ws.add_data_validation(dv_type)
dv_domain.add(f"B2:B1000")
dv_type.add(f"C2:C1000")

# ── ชีต Domains ──────────────────────────────────────────────
wd = wb.create_sheet("Domains")
for i, title in enumerate(["id", "name_en", "name_th", "file"], start=1):
    c = wd.cell(row=1, column=i, value=title)
    c.fill = HEAD_FILL
    c.font = HEAD_FONT
for r, d in enumerate(DOMAINS, start=2):
    wd.cell(row=r, column=1, value=d["id"])
    wd.cell(row=r, column=2, value=d["name"]["en"])
    wd.cell(row=r, column=3, value=d["name"]["th"])
    wd.cell(row=r, column=4, value=d["file"])
wd.column_dimensions["A"].width = 6
wd.column_dimensions["B"].width = 44
wd.column_dimensions["C"].width = 46
wd.column_dimensions["D"].width = 40
wd.freeze_panes = "A2"

# ── ชีต README ───────────────────────────────────────────────
wr = wb.create_sheet("README")
lines = [
    ("วิธีใช้เทมเพลตนี้", True),
    ("", False),
    ("1. กรอกโจทย์ในชีต Questions แถวละหนึ่งข้อ", False),
    ("2. บันทึกไฟล์ แล้วเปิด tools/import.html ในเว็บ ลากไฟล์นี้มาวาง", False),
    ("3. ตรวจรายการปัญหาที่ระบบขึ้นให้ แก้ใน Excel แล้ววางใหม่จนไม่มีข้อผิดพลาด", False),
    ("4. เลือกวิธีรวมกับของเดิม แล้วกดดาวน์โหลด JSON ไป commit ลงโฟลเดอร์ data/", False),
    ("", False),
    ("คอลัมน์ที่จำเป็น (หัวตารางพื้นแดง)", True),
    ("  domain   หมายเลขโดเมน ดูรายการในชีต Domains", False),
    ("  q_en     คำถามภาษาอังกฤษ", False),
    ("  answer   เฉลย เช่น B หรือ B,D ถ้าตอบได้หลายข้อ", False),
    ("", False),
    ("กติกาที่ควรรู้", True),
    ("  • เว้น id ว่างไว้ได้ ระบบจะรันเลขต่อจากของเดิมให้เอง", False),
    ("  • ปล่อยคอลัมน์ e_* และ f_* ว่างได้ถ้ามีตัวเลือกไม่ครบ 6 ข้อ", False),
    ("  • answer ต้องตรงกับตัวอักษรของตัวเลือกที่มีจริงในแถวนั้น", False),
    ("  • type เว้นว่าง = single ระบบจะเปลี่ยนเป็น multi ให้เองถ้า answer มีมากกว่าหนึ่งตัว", False),
    ("  • คอลัมน์ _th เว้นว่างได้ ระบบจะแสดงภาษาอังกฤษแทนและนับเข้าป้าย TH x/y", False),
    ("  • ไม่มี explain_en จะขึ้นเป็นคำเตือน แต่ยังนำเข้าได้", False),
    ("  • insight_en / insight_th ไม่บังคับเลย — ปกติกรอกแค่ insight_th อย่างเดียวก็พอ", False),
    ("    ระบบจะโชว์เนื้อหานี้ให้ทั้งตอนสลับเป็นโหมด EN และ TH โดยอัตโนมัติ (fallback แบบเดียวกับคำอธิบาย)", False),
    ("  • ในเซลล์ Excel กด Alt+Enter เพื่อขึ้นบรรทัดใหม่ได้ ระบบจะคงการขึ้นบรรทัดนั้นไว้ตอนแสดงผล", False),
    ("  • โดเมนที่ยังไม่มีโจทย์เลยจะไม่โผล่ในหน้าเลือกโดเมนของเว็บ — พอมีแถวแรกของโดเมนนั้นเข้าไปก็จะขึ้นเอง", False),
    ("", False),
    ("คำอธิบายรายคอลัมน์", True),
]
for name, _, note in HEADERS:
    lines.append((f"  {name:<12} {note}", False))

for r, (text, bold) in enumerate(lines, start=1):
    c = wr.cell(row=r, column=1, value=text)
    if bold:
        c.font = Font(bold=True, size=11)
wr.column_dimensions["A"].width = 110

out = ROOT / "tools" / "template.xlsx"
wb.save(out)
# ใช้ ASCII ล้วนเพื่อไม่ให้ console บน Windows (cp1252) พัง
print(f"written: {out.relative_to(ROOT)}")
